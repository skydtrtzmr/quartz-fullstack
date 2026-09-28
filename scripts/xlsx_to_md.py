import openpyxl
from pathlib import Path

src = r"d:\CodeProjects\quartz-fullstack\知识图谱功能升级研发项目-任务清单.xlsx"
out = Path(r"d:\CodeProjects\quartz-fullstack\知识图谱功能升级研发项目-任务清单.md")

wb = openpyxl.load_workbook(src, data_only=True)
parts = []

for ws in wb.worksheets:
    parts.append(f"# {ws.title}\n")
    merged = {}
    for rng in ws.merged_cells.ranges:
        val = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged[(r, c)] = val
    rows = []
    for r in range(1, ws.max_row + 1):
        row = []
        for c in range(1, ws.max_column + 1):
            v = merged.get((r, c))
            if v is None:
                v = ws.cell(r, c).value
            v = "" if v is None else str(v)
            v = v.replace("|", "\\|").replace("\r", "").replace("\n", "<br>")
            row.append(v)
        rows.append(row)
    if not rows:
        continue
    ncol = max(len(r) for r in rows)
    header = rows[0] + [""] * (ncol - len(rows[0]))
    parts.append("| " + " | ".join(header) + " |")
    parts.append("| " + " | ".join(["---"] * ncol) + " |")
    for row in rows[1:]:
        row = row + [""] * (ncol - len(row))
        parts.append("| " + " | ".join(row) + " |")
    parts.append("")

out.write_text("\n".join(parts), encoding="utf-8")
print("OK", out, "sheets:", wb.sheetnames)
